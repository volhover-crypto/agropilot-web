# tests/test_knowledge.py -- О1 (§41): конвейер документов, чанкинг,
# кодировки, идемпотентность, валидация цитат.
# Конвейер гоняется на sqlite+aiosqlite с FakeEmbedder + MemoryStore
# (без Qdrant/fastembed -- тяжёлые зависимости только в проде).

import asyncio
import os

os.environ["EMBED_PROVIDER"] = "fake"      # тесты -- без heavy-зависимостей
os.environ["QDRANT_URL"] = ""              # MemoryStore

from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine  # noqa: E402

from backend.knowledge.chunking import chunk_text  # noqa: E402
from backend.knowledge.embedder import FakeEmbedder  # noqa: E402
from backend.knowledge.ingest import IngestError, extract_text, run_pipeline  # noqa: E402
from backend.knowledge.models import Base, KnowledgeBase, KnowledgeDoc, KnowledgeChunk  # noqa: E402
from backend.knowledge.retrieval import validate_citations  # noqa: E402
from backend.knowledge.store import MemoryStore  # noqa: E402

LONG_TEXT = ("Капельный полив виноградника требует нормирования воды. " * 60)


def _run_with_session(coro_fn):
    """Создаёт sqlite-БД и сессию, прогоняет coro_fn(session) в ОДНОМ
    event loop (aiosqlite-соединение привязано к циклу)."""
    async def wrap():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        maker = async_sessionmaker(engine, expire_on_commit=False)
        async with maker() as session:
            return await coro_fn(session)

    return asyncio.run(wrap())


# ---------- конвейер (sqlite) ----------

async def _new_doc(session, title):
    kb = KnowledgeBase(title="Тестовая KB")
    session.add(kb)
    await session.flush()
    doc = KnowledgeDoc(kb_id=kb.id, title=title)
    session.add(doc)
    await session.commit()
    await session.refresh(kb)
    await session.refresh(doc)
    return kb, doc


def test_pipeline_ready_and_corpus_version():
    async def scenario(session):
        kb, doc = await _new_doc(session, "полив.txt")
        store, emb = MemoryStore(), FakeEmbedder()
        doc = await run_pipeline(session, doc, LONG_TEXT.encode("utf-8"),
                                 store=store, embedder=emb)
        from sqlalchemy import select as sa_select

        kb2 = await session.get(KnowledgeBase, kb.id)
        chunks = (await session.execute(sa_select(KnowledgeChunk))).scalars().all()
        return doc, kb2, chunks

    doc, kb2, chunks = _run_with_session(scenario)
    assert doc.status == "ready" and doc.charset == "utf-8"
    assert kb2.corpus_version == 1 and kb2.doc_count == 1
    assert len(chunks) >= 3
    assert all(c.qdrant_point_id for c in chunks)


def test_pipeline_idempotent_no_duplicates():
    async def scenario(session):
        from sqlalchemy import select as sa_select

        kb, doc = await _new_doc(session, "полив.txt")
        store, emb = MemoryStore(), FakeEmbedder()

        async def run():
            return await run_pipeline(session, doc, LONG_TEXT.encode("utf-8"),
                                      store=store, embedder=emb)

        await run()
        n1 = len((await session.execute(sa_select(KnowledgeChunk))).scalars().all())
        doc2 = await run()  # повторная индексация того же doc_id
        n2 = len((await session.execute(sa_select(KnowledgeChunk))).scalars().all())
        pts = [1 for _, (v, pl) in store._points.items() if pl.get("doc_id") == doc.id]
        return doc2, n1, n2, len(pts)

    doc2, n1, n2, npts = _run_with_session(scenario)
    assert doc2.status == "ready"
    assert n1 == n2 and n1 > 0  # дублей чанков нет
    assert npts == n1  # и точек в хранилище ровно столько же


def test_pipeline_encoding_fail_keeps_corpus():
    async def scenario(session):
        kb, doc = await _new_doc(session, "битый.txt")
        doc = await run_pipeline(session, doc, b"\x98" * 100,
                                 store=MemoryStore(), embedder=FakeEmbedder())
        kb2 = await session.get(KnowledgeBase, kb.id)
        return doc, kb2

    doc, kb2 = _run_with_session(scenario)
    assert doc.status == "failed" and doc.fail_reason == "encoding"
    assert kb2.corpus_version == 0  # корпус не изменился


# ---------- чанкинг ----------

def test_chunking_covers_text_with_overlap():
    chunks = chunk_text(LONG_TEXT, size=500, overlap=0.15)
    assert len(chunks) >= 3
    # полное покрытие и монотонные диапазоны
    assert chunks[0]["start"] == 0
    for a, b in zip(chunks, chunks[1:]):
        assert b["start"] > a["start"]
        assert b["start"] < a["end"]  # перекрытие ~15%
    assert [c["ord"] for c in chunks] == list(range(len(chunks)))
    c0 = chunks[0]
    assert c0["text"].strip() == LONG_TEXT[c0["start"]:c0["end"]].strip()


def test_chunking_empty_text():
    assert chunk_text("") == []
    assert chunk_text("   \n  ") == []


# ---------- кодировки и извлечение ----------

def test_decode_utf8_and_cp1251_fallback():
    text, enc = extract_text("осень.txt", "Норма полива 120 м3".encode("utf-8"))
    assert enc == "utf-8" and "Норма" in text
    text, enc = extract_text("осень.txt", "Норма полива 120 м3".encode("cp1251"))
    assert enc == "cp1251" and "Норма" in text


def test_decode_undecodable_fails_encoding():
    try:
        extract_text("битый.txt", b"\x98" * 100)
        raised = False
    except IngestError as e:
        raised = e.reason == "encoding"
    assert raised


def test_extract_docx_and_xlsx():
    import io

    from docx import Document

    d = Document()
    d.add_paragraph("Поливная норма для сада — 600 м3/га.")
    buf = io.BytesIO()
    d.save(buf)
    text, _ = extract_text("нормы.docx", buf.getvalue())
    assert "600" in text

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Культура", "Норма"])
    ws.append(["Виноград", "550"])
    buf2 = io.BytesIO()
    wb.save(buf2)
    text2, _ = extract_text("нормы.xlsx", buf2.getvalue())
    assert "Виноград" in text2 and "550" in text2


def test_extract_csv_cp1251():
    data = "культура;норма\nВиноград;550\n".encode("cp1251")
    text, enc = extract_text("нормы.csv", data)
    assert enc == "cp1251" and "Виноград" in text


def test_pdf_without_text_layer_fails_no_text():
    """Фикс-паттерн Octop: OCR-отказ («нет изображения») -> failed(no_text),
    НЕ ready с пустым корпусом."""
    import io

    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    try:
        extract_text("скан.pdf", buf.getvalue())
        raised = False
    except IngestError as e:
        raised = e.reason == "no_text"
    assert raised


def test_broken_pdf_fails_parse():
    try:
        extract_text("битый.pdf", b"not a pdf at all")
        raised = False
    except IngestError as e:
        raised = e.reason in ("parse", "no_text")
    assert raised


# ---------- memory store + цитаты ----------

def test_memory_store_search_and_delete_by_doc():
    emb, store = FakeEmbedder(), MemoryStore()
    texts = ["капельный полив виноградника норма", "финансовый отчёт по сделкам"]
    vecs = emb.embed(texts)
    store.ensure_collection(emb.dim())
    store.upsert([
        {"id": "doc1:c1", "vector": vecs[0], "payload": {"doc_id": 1, "kb_id": 1, "chunk_id": 11, "ord": 0}},
        {"id": "doc2:c2", "vector": vecs[1], "payload": {"doc_id": 2, "kb_id": 1, "chunk_id": 22, "ord": 0}},
    ])
    hits = store.search(emb.embed(["капельный полив"])[0], limit=2)
    assert hits[0]["payload"]["doc_id"] == 1 and hits[0]["score"] > 0.3
    assert store.delete_by_doc(1) == 1
    hits2 = store.search(emb.embed(["капельный полив"])[0], limit=2)
    assert all(h["payload"]["doc_id"] != 1 for h in hits2)


def test_validate_citations():
    assert validate_citations("Норма [1], а здесь [2].", 3) == [1, 2]
    assert validate_citations("ссылка [9] вне диапазона", 3) == []   # unverified
    assert validate_citations("без ссылок вообще", 3) == []          # unverified
    assert validate_citations("", 3) == []


def test_point_ids_are_uuid_for_qdrant():
    """Qdrant принимает только UUID/int ID: qdrant_point_id обязан быть
    валидным UUID (детерминированный uuid5, стабильный при переиндексации)."""
    import uuid

    async def scenario(session):
        kb, doc = await _new_doc(session, "полив.txt")
        doc = await run_pipeline(session, doc, LONG_TEXT.encode("utf-8"),
                                 store=MemoryStore(), embedder=FakeEmbedder())
        from sqlalchemy import select as sa_select

        chunks = (await session.execute(sa_select(KnowledgeChunk))).scalars().all()
        return [c.qdrant_point_id for c in chunks]

    pids = _run_with_session(scenario)
    assert pids and all(str(uuid.UUID(p)) == p for p in pids)
